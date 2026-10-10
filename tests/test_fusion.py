"""Synthetic integration cases; these do not measure extraction quality."""
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from bmw_pm.contracts import PolicyVersion, SurveyObservation, SourceLocator
from bmw_pm.ingestion import FeedbackWorkbookAdapter
from bmw_pm.ingestion.feedback import EXPECTED_HEADERS
from bmw_pm.intelligence import BaselineExtractionProvider
from bmw_pm.intelligence.fusion import import_youtube_run, read_fusion_inputs, fuse_candidates, validate_cross_channel_links
from bmw_pm.pipeline import load_json, run_candidate_pipeline
from bmw_pm.storage import Database


def saved_run(root, *, generation="G60", quote="The touchscreen is hard to use.", kind="pain_point"):
    root.mkdir(parents=True)
    (root / "transcripts").mkdir()
    (root / "transcripts/abc.json").write_text(json.dumps([
        {"start": 10, "duration": 4, "text": quote}]), encoding="utf-8")
    source = {"video_id": "abc", "retrieval_status": "retrieved", "title": "Synthetic review",
              "channel": "Synthetic channel", "url": "https://youtube.com/watch?v=abc"}
    (root / "sources.json").write_text(json.dumps([source]), encoding="utf-8")
    observation = {"video_id": "abc", "generation": generation, "topic": "touchscreen",
                   "source_statement": quote, "supporting_excerpt": quote,
                   "timestamp_seconds": 10, "feedback_type": kind}
    (root / "evidence.json").write_text(json.dumps([observation, observation]), encoding="utf-8")
    return root


def database(tmp_path, feedback=True):
    db = Database(tmp_path / "test.db")
    db.initialize()
    if feedback:
        path = tmp_path / "G60.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Feedback_Explorer"
        sheet.append(list(EXPECTED_HEADERS))
        sheet.append(["G60-001", "A", "BMW", "BMW 5", "US", "G60", "ICE", "Negative",
                      "The touchscreen is hard to use.", "controls", "controls", "controls"])
        book.save(path)
        batch = FeedbackWorkbookAdapter().parse(path, model_code="G60")
        db.persist_feedback_ingestion(batch, request_id="fixture")
    return db


def run(db, **kwargs):
    return run_candidate_pipeline(database=db, profile=load_json("configs/g60.json"),
        policy=PolicyVersion.model_validate(load_json("configs/exploratory_equal_v1.json")),
        provider=BaselineExtractionProvider.from_file("configs/baseline_taxonomy_v1.json"), **kwargs)


def controls(result):
    return next(r for r in result["requirements"] if r["title"] == "Controls Usability")


def test_joint_engine_preserves_denominator_and_new_evidence_revision(tmp_path):
    db = database(tmp_path)
    initial = run(db, youtube_runs=[])
    before = controls(initial)
    path = saved_run(tmp_path / "youtube", kind="positive", quote="The touchscreen is easy to use.")
    joint = run(db, youtube_runs=[path])
    card = controls(joint)
    assert card["support_count"] == before["support_count"] == 1
    assert card["denominator_feedback_records"] == 1
    assert len(card["cross_channel_evidence"]) == 1
    assert card["cross_channel_evidence"][0]["role"] == "preserves"
    assert any("Mixed evidence" in s for s in card["evidence_synthesis"])
    assert card["revision"] == before["revision"] + 1
    assert card["feasibility"] == "unknown"
    assert card["implementation_cost"] is None
    assert joint["run"]["corpus_manifest_hash"] != initial["run"]["corpus_manifest_hash"]
    repeated = run(db, youtube_runs=[path])
    assert repeated["run"]["corpus_manifest_hash"] == joint["run"]["corpus_manifest_hash"]
    repeat = controls(repeated)
    assert repeat["revision"] == card["revision"]
    assert len(read_fusion_inputs(db)["youtube"]) == 1
    replay = db.replay_requirement(card["requirement_id"])
    assert len(replay["revisions"]) == 2 and replay["audit_valid"]
    assert replay["revisions"][0]["cross_channel_evidence"] == []
    assert validate_cross_channel_links(db, [db.latest_requirement(card["requirement_id"])]) == []
    with db.connect() as connection:
        connection.execute("UPDATE youtube_bundles SET payload_json='{}'")
        connection.commit()
    assert validate_cross_channel_links(db, [db.latest_requirement(card["requirement_id"])])


def test_youtube_only_need_and_scope_filter(tmp_path):
    db = database(tmp_path, feedback=False)
    path = saved_run(tmp_path / "youtube")
    result = run(db, youtube_runs=[path])
    card = controls(result)
    assert card["support_count"] == 0
    assert card["cross_channel_evidence"][0]["applicability"] == "scope_unconfirmed"
    score = next(s for s in result["scores"] if s["requirement_id"] == card["requirement_id"])
    assert score["criterion_inputs"]["evidence_backed_reach"]["value"] is None
    foreign = saved_run(tmp_path / "foreign", generation="G70")
    other = database(tmp_path / "other", feedback=False)
    with pytest.raises(ValueError, match="no eligible"):
        run(other, youtube_runs=[foreign])


@pytest.mark.parametrize("mutation", ["quote", "timestamp", "path", "missing_transcript"])
def test_rejects_unresolvable_youtube_evidence(tmp_path, mutation):
    db = database(tmp_path, feedback=False)
    path = saved_run(tmp_path / "youtube")
    rows = json.loads((path / "evidence.json").read_text())
    if mutation == "quote":
        rows[0]["supporting_excerpt"] = "Invented claim"
    elif mutation == "timestamp":
        rows[0]["timestamp_seconds"] = 50
    elif mutation == "path":
        rows[0]["video_id"] = "../../secret"
    else:
        (path / "transcripts/abc.json").unlink()
    (path / "evidence.json").write_text(json.dumps(rows[:1]))
    report = import_youtube_run(db, path)
    assert report["accepted"] == 0 and len(report["rejected"]) == 1
    assert read_fusion_inputs(db)["youtube"] == []


def test_studies_and_family_sales_remain_scoped_context(tmp_path):
    db = database(tmp_path)
    base = run(db)
    observation = SurveyObservation(survey_id="study", source_id="source", market="EU", model="G60",
        instrument_id="test", category_raw="touchscreen", statistic_type="mean", value=7,
        year=2025, mapping_version="test", cell_locators=[SourceLocator(source_id="source",
            source_hash="a" * 64, locator_type="sheet_range", sheet="EU", cell_range="E4")], limitations=[])
    imported = {"row_key": "sale", "archive_sha256": "b" * 64, "table_name": "sales_observation",
                "payload": {"model": "G60_G68", "market": "CN", "volume": 1000}}
    candidates = [db.latest_requirement(r["requirement_id"]) for r in base["requirements"]]
    fused = fuse_candidates(candidates=candidates, inputs={"youtube": [], "studies": [observation], "imported": [imported]},
        taxonomy=load_json("configs/baseline_taxonomy_v1.json"), eligible=db.feedback_records(model_code="G60"),
        processed_ids=set(), profile=load_json("configs/g60.json"), created_by="test")
    card = next(c for c in fused if c.title == "Controls Usability")
    assert {l.channel for l in card.cross_channel_evidence} == {"bmw_study", "imported_context"}
    assert all(l.role == "context" for l in card.cross_channel_evidence)
    assert card.support_count == 1


def test_empty_blocked_run_does_not_generate_findings(tmp_path):
    db = database(tmp_path)
    path = tmp_path / "empty"
    path.mkdir()
    (path / "evidence.json").write_text("[]")
    (path / "sources.json").write_text(json.dumps([{"video_id": "abc", "retrieval_status": "skipped_ip_blocked"}]))
    result = run(db, youtube_runs=[path])
    assert result["youtube_imports"][0]["accepted"] == 0
    assert result["youtube_imports"][0]["readable_sources"] == 0
    assert controls(result)["cross_channel_evidence"] == []


def test_workbench_renders_joint_assessment_and_citations(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    app_path = Path("ui/app.py").resolve()
    db = database(tmp_path)
    run(db, youtube_runs=[saved_run(tmp_path / "youtube")])
    monkeypatch.delenv("BMW_PM_HOSTED", raising=False)
    monkeypatch.delenv("RENDER", raising=False)
    monkeypatch.chdir(tmp_path)
    app = AppTest.from_file(str(app_path)).run(timeout=30)
    next(t for t in app.text_input if t.label == "SQLite database").set_value(str(db.path)).run(timeout=30)
    assert not app.exception
    assert any(s.value == "Combined evidence assessment" for s in app.subheader)
    assert any(s.label == "Cross-channel evidence" for s in app.selectbox)


def test_inapplicable_charging_and_unmatched_topics_are_reported(tmp_path):
    db = database(tmp_path, feedback=False)
    path = saved_run(tmp_path / "youtube", generation="F70", quote="Charging is slow.")
    profile = {**load_json("configs/g60.json"), "model_family": "F70", "powertrains": ["ICE"]}
    result = run_candidate_pipeline(database=db, profile=profile,
        policy=PolicyVersion.model_validate(load_json("configs/exploratory_equal_v1.json")),
        provider=BaselineExtractionProvider.from_file("configs/baseline_taxonomy_v1.json"),
        youtube_runs=[path])
    assert not any(r["title"] == "Charging And Range" for r in result["requirements"])
    # The metadata topic still mentions touchscreen; generation scope is retained for review.
    assert all(l["applicability"] == "scope_unconfirmed" for r in result["requirements"] for l in r["cross_channel_evidence"])
    rows = json.loads((path / "evidence.json").read_text())
    for row in rows:
        row["topic"] = "charging"
    (path / "evidence.json").write_text(json.dumps(rows))
    # Use a fresh DB because existing imported findings are intentionally immutable.
    fresh = database(tmp_path / "fresh", feedback=False)
    result = run_candidate_pipeline(database=fresh, profile=profile,
        policy=PolicyVersion.model_validate(load_json("configs/exploratory_equal_v1.json")),
        provider=BaselineExtractionProvider.from_file("configs/baseline_taxonomy_v1.json"), youtube_runs=[path])
    assert result["requirements"] == []
    assert len(result["unlinked_youtube_findings"]) == 1


def sales_book(path, bad_formula=False):
    book = Workbook()
    book.remove(book.active)
    for model in ("F70", "G60_G68", "G70"):
        sheet = book.create_sheet(model)
        sheet.append(["market", "market_code", "volume_2024", "volume_2025", "volume_2030"])
        for label, region in (("Europe", "EU"), ("China", "CN"), ("USA", "US"), ("Rest of World", "RoW")):
            sheet.append([label, region, 0, 100, 120])
    if bad_formula:
        book["G60_G68"]["C2"] = "=1+1"
    book.save(path)


def test_supplied_sales_context_join_idempotency_and_provenance(tmp_path):
    from bmw_pm.ingestion.context import ingest_provided_context, parse_sales
    from bmw_pm.validation import validate_database
    db = database(tmp_path)
    sales_book(tmp_path / "sales_volumes.xlsx")
    source, observations = parse_sales(tmp_path / "sales_volumes.xlsx")
    assert len(observations) == 36
    assert any(o.value == 0 for o in observations)
    assert any("planning scenario" in " ".join(o.limitations) for o in observations if o.year == 2030)
    assert ingest_provided_context(db, tmp_path)["inserted"] == 36
    assert ingest_provided_context(db, tmp_path)["inserted"] == 0
    joint = run(db, youtube_runs=[])
    links = [l for l in controls(joint)["cross_channel_evidence"] if l["channel"] == "bmw_sales"]
    assert len(links) == 12
    assert all(l["payload"]["model_codes"] == ["G60", "G68"] for l in links)
    assert all(l["role"] == "context" for l in links)
    assert controls(joint)["support_count"] == 1
    assert validate_cross_channel_links(db, [db.latest_requirement(r["requirement_id"]) for r in joint["requirements"]]) == []
    check = next(c for c in validate_database(db)["checks"] if c["check"] == "provided_context_reconciliation")
    assert check["passed"]


def test_sales_formulas_without_cached_values_are_not_evaluated(tmp_path):
    from bmw_pm.ingestion.context import parse_sales
    from bmw_pm.ingestion.feedback import IngestionError
    path = tmp_path / "sales.xlsx"
    sales_book(path, bad_formula=True)
    with pytest.raises(IngestionError, match="formulas are never evaluated"):
        parse_sales(path)


def test_equipment_blank_page_is_retained_as_unknown_not_no_feature(tmp_path):
    PdfWriter = pytest.importorskip("pypdf").PdfWriter
    from bmw_pm.ingestion.context import parse_equipment
    path = tmp_path / "G60_OptionList.PDF"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(path)
    source, pages = parse_equipment(path, "G60")
    assert len(pages) == 1 and pages[0].market == "DE"
    assert pages[0].locator.page == 1
    assert pages[0].raw_payload["extraction_status"] == "needs_ocr_or_visual_review"
    assert pages[0].locator.source_hash == source.sha256
    assert "does not establish" in " ".join(pages[0].limitations)


def test_cli_json_survives_windows_legacy_stdout(monkeypatch):
    import io
    from bmw_pm import cli
    raw = io.BytesIO()
    stdout = io.TextIOWrapper(raw, encoding="cp1252")
    with monkeypatch.context() as patch:
        patch.setattr(cli.sys, "stdout", stdout)
        cli._json_print({"quote": "■ Klimatisierung"})
        stdout.flush()
    assert json.loads(raw.getvalue().decode("cp1252"))["quote"] == "■ Klimatisierung"

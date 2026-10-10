from __future__ import annotations

import mimetypes
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..contracts import (
    QuarantineFinding,
    SourceDocument,
    SourceLocator,
    StudyIngestionResult,
    SurveyObservation,
)
from ..utils import sha256_file, stable_id, utc_now
from .feedback import IngestionError


US_RESPONSE_LABELS = (
    "I Hate It",
    "A Failure",
    "Unsatisfactory",
    "Satisfactory",
    "Excellent",
    "Delightful",
    "I Love It",
)


class CustomerStudyWorkbookAdapter:
    parser_version = "customer-study-xlsx-v1"

    def parse(self, path: str | Path, *, strict_expected_counts: bool = True) -> StudyIngestionResult:
        file_path = Path(path).resolve()
        source_hash = sha256_file(file_path)
        source_id = stable_id("source", source_hash)
        source = SourceDocument(
            source_id=source_id,
            local_or_web="local",
            original_name_or_url=file_path.name,
            sha256=source_hash,
            byte_count=file_path.stat().st_size,
            mime_type=mimetypes.guess_type(file_path.name)[0]
            or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            acquired_at=utc_now(),
            source_type="supplied_customer_study_workbook",
            access_scope="local_restricted",
            origin_family_id=source_id,
            parser_version=self.parser_version,
            rights_notes="Supplied aggregate challenge material; respondent counts and scale endpoints are absent.",
        )
        # Study parsing uses random cell access; streaming mode repeatedly reparses XML.
        # The supplied aggregate workbook is small. Formulas remain unevaluated.
        workbook = load_workbook(file_path, read_only=False, data_only=False)
        observations: list[SurveyObservation] = []
        quarantined: list[QuarantineFinding] = []
        try:
            required = {"CN_EU_2025", "US_2025"}
            if not required.issubset(workbook.sheetnames):
                raise IngestionError("customer-study workbook is missing a required sheet")
            observations.extend(
                self._parse_cn_eu(workbook["CN_EU_2025"], source_id, source_hash, quarantined)
            )
            observations.extend(
                self._parse_us(workbook["US_2025"], source_id, source_hash, quarantined)
            )
        finally:
            workbook.close()
        if strict_expected_counts:
            means = sum(item.statistic_type == "mean" for item in observations)
            distributions = sum(
                item.statistic_type == "categorical_distribution" for item in observations
            )
            if means != 185 or distributions != 206:
                raise IngestionError(
                    f"study reconciliation failed: {means} means and {distributions} distributions"
                )
        return StudyIngestionResult(
            source=source, observations=observations, quarantined=quarantined
        )

    def _parse_cn_eu(
        self,
        sheet: Any,
        source_id: str,
        source_hash: str,
        quarantined: list[QuarantineFinding],
    ) -> list[SurveyObservation]:
        scopes = {
            2: ("CN", "G68"),
            3: ("CN", "G70"),
            4: ("EU_UNRESOLVED_EUROPE_SCOPE", "F70"),
            5: ("EU_UNRESOLVED_EUROPE_SCOPE", "G60"),
            6: ("EU_UNRESOLVED_EUROPE_SCOPE", "G70"),
        }
        results: list[SurveyObservation] = []
        for category_row in range(3, 76, 2):
            category = sheet.cell(category_row, 1).value
            mean_label = sheet.cell(category_row + 1, 1).value
            if not category or mean_label != "Mean":
                quarantined.append(
                    QuarantineFinding(
                        source_id=source_id,
                        source_record_id=f"CN_EU_2025-row-{category_row}",
                        row_numbers=[category_row, category_row + 1],
                        reason="Expected category/Mean row pair was not present",
                    )
                )
                continue
            for column, (market, model) in scopes.items():
                value = sheet.cell(category_row + 1, column).value
                if not isinstance(value, (int, float)):
                    quarantined.append(
                        QuarantineFinding(
                            source_id=source_id,
                            source_record_id=f"CN_EU_2025-{category_row}-{column}",
                            row_numbers=[category_row + 1],
                            reason="Mean cell was not numeric",
                        )
                    )
                    continue
                cell = sheet.cell(category_row + 1, column).coordinate
                results.append(
                    SurveyObservation(
                        survey_id=stable_id(
                            "survey", source_hash, "CN_EU_2025", str(category), market, model
                        ),
                        source_id=source_id,
                        market=market,
                        model=model,
                        instrument_id="CN_EU_2025",
                        category_raw=str(category),
                        statistic_type="mean",
                        value=float(value),
                        year=2025,
                        mapping_version="study-layout-v1",
                        cell_locators=[
                            SourceLocator(
                                source_id=source_id,
                                source_hash=source_hash,
                                locator_type="sheet_range",
                                sheet="CN_EU_2025",
                                cell_range=cell,
                                extracted_value=str(value),
                            )
                        ],
                        limitations=[
                            "Scale endpoints are not supplied.",
                            "Respondent count and weighting method are not supplied.",
                            "The workbook label EU is retained as an unresolved Europe reporting scope.",
                        ],
                    )
                )
        return results

    def _parse_us(
        self,
        sheet: Any,
        source_id: str,
        source_hash: str,
        quarantined: list[QuarantineFinding],
    ) -> list[SurveyObservation]:
        results: list[SurveyObservation] = []
        row = 1
        while row <= sheet.max_row:
            category = sheet.cell(row, 1).value
            if (
                not category
                or row + 1 > sheet.max_row
                or sheet.cell(row + 1, 1).value != "Sample total"
            ):
                row += 1
                continue
            response_rows: list[int] = []
            cursor = row + 2
            while cursor <= sheet.max_row and len(response_rows) < len(US_RESPONSE_LABELS):
                label = sheet.cell(cursor, 1).value
                if label is None:
                    cursor += 1
                    continue
                response_rows.append(cursor)
                cursor += 1
            labels = tuple(sheet.cell(item, 1).value for item in response_rows)
            if labels != US_RESPONSE_LABELS:
                quarantined.append(
                    QuarantineFinding(
                        source_id=source_id,
                        source_record_id=f"US_2025-row-{row}",
                        row_numbers=list(range(row, min(cursor, sheet.max_row + 1))),
                        reason="Category block did not contain the expected seven response labels",
                    )
                )
                row += 1
                continue
            for column, model in ((3, "G60"), (4, "G70")):
                values = [sheet.cell(item, column).value for item in response_rows]
                if not all(isinstance(value, (int, float)) for value in values):
                    quarantined.append(
                        QuarantineFinding(
                            source_id=source_id,
                            source_record_id=f"US_2025-{row}-{model}",
                            row_numbers=response_rows,
                            reason="Distribution block contained a non-numeric response share",
                        )
                    )
                    continue
                distribution = {
                    label: float(value) for label, value in zip(US_RESPONSE_LABELS, values)
                }
                if abs(sum(distribution.values()) - 1.0) > 0.005:
                    quarantined.append(
                        QuarantineFinding(
                            source_id=source_id,
                            source_record_id=f"US_2025-{row}-{model}",
                            row_numbers=response_rows,
                            reason="Distribution shares did not sum to one within the declared 0.005 parser tolerance",
                        )
                    )
                    continue
                column_letter = "C" if column == 3 else "D"
                results.append(
                    SurveyObservation(
                        survey_id=stable_id(
                            "survey", source_hash, "US_2025", str(category), model
                        ),
                        source_id=source_id,
                        market="US",
                        model=model,
                        instrument_id="US_2025",
                        category_raw=str(category),
                        statistic_type="categorical_distribution",
                        distribution=distribution,
                        year=2025,
                        mapping_version="study-layout-v1",
                        cell_locators=[
                            SourceLocator(
                                source_id=source_id,
                                source_hash=source_hash,
                                locator_type="sheet_range",
                                sheet="US_2025",
                                cell_range=(
                                    f"{column_letter}{response_rows[0]}:"
                                    f"{column_letter}{response_rows[-1]}"
                                ),
                                extracted_value=str(distribution),
                            )
                        ],
                        limitations=[
                            "Sample total of 1 is a proportion base, not one respondent.",
                            "Respondent count and weighting method are not supplied.",
                            "The ordinal distribution is retained without an assumed equal-interval mean.",
                        ],
                    )
                )
            row = cursor
        return results

from __future__ import annotations

import mimetypes
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..contracts import (
    FeedbackIngestionResult,
    FeedbackRecord,
    LegacyAnnotation,
    QuarantineFinding,
    SourceDocument,
    SourceLocator,
    TransformationLog,
)
from ..utils import (
    normalize_whitespace,
    redact_pii_same_length,
    sha256_bytes,
    sha256_file,
    stable_id,
    utc_now,
)


class IngestionError(ValueError):
    pass


EXPECTED_HEADERS = (
    "ID",
    "Source",
    "Brand",
    "Sales Description",
    "Country",
    "Derivate (e-code)",
    "Engine Type",
    "Feedback Type",
    "Customer Feedback",
    "Vfc level2 3 Name",
    "Vfc level2 Name",
    "Vfc level3 Name",
)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    return str(value)


class FeedbackWorkbookAdapter:
    parser_version = "feedback-xlsx-v1"

    def __init__(
        self,
        *,
        powertrain_aliases: dict[str, str] | None = None,
        alias_assumptions: dict[str, str] | None = None,
    ) -> None:
        self.powertrain_aliases = powertrain_aliases or {}
        self.alias_assumptions = alias_assumptions or {}

    def parse(self, path: str | Path, *, model_code: str) -> FeedbackIngestionResult:
        file_path = Path(path).resolve()
        source_hash = sha256_file(file_path)
        source_id = stable_id("source", source_hash)
        source = SourceDocument(
            source_id=source_id,
            local_or_web="local",
            original_name_or_url=file_path.name,
            sha256=source_hash,
            byte_count=file_path.stat().st_size,
            mime_type=mimetypes.guess_type(file_path.name)[0]
            or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            acquired_at=utc_now(),
            source_type="supplied_anonymized_feedback_workbook",
            access_scope="local_private",
            origin_family_id=source_id,
            parser_version=self.parser_version,
            rights_notes="Supplied challenge material; retain locally and do not commit raw content.",
        )
        workbook = load_workbook(file_path, read_only=True, data_only=False)
        try:
            if "Feedback_Explorer" not in workbook.sheetnames:
                raise IngestionError("missing required Feedback_Explorer sheet")
            sheet = workbook["Feedback_Explorer"]
            rows = sheet.iter_rows(values_only=True)
            headers = tuple(next(rows, ()))
            if headers[: len(EXPECTED_HEADERS)] != EXPECTED_HEADERS:
                raise IngestionError(
                    "unexpected feedback headers; ambiguous mappings are not guessed"
                )
            grouped: dict[str, list[tuple[int, tuple[Any, ...]]]] = defaultdict(list)
            input_rows = 0
            for excel_row, row in enumerate(rows, start=2):
                if not any(value is not None for value in row):
                    continue
                input_rows += 1
                record_key = _text(row[0])
                if not record_key:
                    grouped[f"__missing_id_row_{excel_row}"].append((excel_row, row))
                else:
                    grouped[record_key].append((excel_row, row))
        finally:
            workbook.close()

        records: list[FeedbackRecord] = []
        annotations: list[LegacyAnnotation] = []
        transformations: list[TransformationLog] = []
        quarantined: list[QuarantineFinding] = []
        for source_record_id, items in grouped.items():
            row_numbers = [number for number, _ in items]
            if source_record_id.startswith("__missing_id_row_"):
                quarantined.append(
                    QuarantineFinding(
                        source_id=source_id,
                        source_record_id=source_record_id,
                        row_numbers=row_numbers,
                        reason="Missing feedback ID",
                    )
                )
                continue
            rows_for_id = [row for _, row in items]
            text_values = {_text(row[8]) for row in rows_for_id}
            metadata_values = {tuple(_text(value) for value in row[1:7]) for row in rows_for_id}
            if len(text_values) != 1 or None in text_values:
                quarantined.append(
                    QuarantineFinding(
                        source_id=source_id,
                        source_record_id=source_record_id,
                        row_numbers=row_numbers,
                        reason="Conflicting or missing feedback text within one source ID",
                    )
                )
                continue
            if len(metadata_values) != 1:
                quarantined.append(
                    QuarantineFinding(
                        source_id=source_id,
                        source_record_id=source_record_id,
                        row_numbers=row_numbers,
                        reason="Conflicting source/model/country/powertrain metadata within one source ID",
                    )
                )
                continue

            raw_text = next(iter(text_values)) or ""
            record_id = stable_id("feedback", source_hash, source_record_id)
            cleaned_text, span_map = normalize_whitespace(raw_text)
            redacted_text, redaction_counts = redact_pii_same_length(cleaned_text)
            normalize_id = stable_id("transform", record_id, "normalize_whitespace_v1")
            redact_id = stable_id("transform", record_id, "pii_redaction_v1")
            transformations.extend(
                [
                    TransformationLog(
                        transformation_id=normalize_id,
                        record_id=record_id,
                        operation="normalize_whitespace",
                        operation_version="1",
                        input_hash=sha256_bytes(raw_text.encode("utf-8")),
                        output_hash=sha256_bytes(cleaned_text.encode("utf-8")),
                        span_map=span_map,
                        details={"reversible_via": "source locator plus span map"},
                    ),
                    TransformationLog(
                        transformation_id=redact_id,
                        record_id=record_id,
                        operation="local_pii_redaction",
                        operation_version="1",
                        input_hash=sha256_bytes(cleaned_text.encode("utf-8")),
                        output_hash=sha256_bytes(redacted_text.encode("utf-8")),
                        details={
                            "same_length_offsets": True,
                            "match_counts": redaction_counts,
                            "warning": "Heuristic redaction requires privacy review before hosted inference.",
                        },
                    ),
                ]
            )
            first = rows_for_id[0]
            powertrain_raw = _text(first[6])
            assumptions: list[str] = []
            if powertrain_raw in self.powertrain_aliases:
                powertrain = self.powertrain_aliases[powertrain_raw]
                mapping_status = "assumed_alias"
                assumptions.append(
                    self.alias_assumptions.get(
                        powertrain_raw,
                        f"{powertrain_raw} was mapped to {powertrain}; confirmation pending.",
                    )
                )
            elif powertrain_raw in {"ICE", "PHEV", "BEV"}:
                powertrain = powertrain_raw
                mapping_status = "exact"
            else:
                powertrain = None
                mapping_status = "unknown"
            row_locators = [
                SourceLocator(
                    source_id=source_id,
                    source_hash=source_hash,
                    locator_type="sheet_range",
                    sheet="Feedback_Explorer",
                    cell_range=f"A{row_number}:L{row_number}",
                    normalized_text_hash=sha256_bytes(cleaned_text.encode("utf-8")),
                )
                for row_number in row_numbers
            ]
            records.append(
                FeedbackRecord(
                    record_id=record_id,
                    source_id=source_id,
                    source_record_id=source_record_id,
                    raw_text_ref=f"{source_id}#Feedback_Explorer!I{row_numbers[0]}",
                    cleaned_text=cleaned_text,
                    redacted_text=redacted_text,
                    raw_country=_text(first[4]),
                    country=_text(first[4]),
                    raw_model=_text(first[3]),
                    model_code=model_code,
                    raw_brand=_text(first[2]),
                    raw_sales_description=_text(first[3]),
                    raw_derivative=_text(first[5]),
                    powertrain_raw=powertrain_raw,
                    powertrain=powertrain,
                    powertrain_mapping_status=mapping_status,
                    source_label=_text(first[1]),
                    row_locators=row_locators,
                    transformation_ids=[normalize_id, redact_id],
                    assumptions=assumptions,
                )
            )
            for row_number, row in items:
                locator = SourceLocator(
                    source_id=source_id,
                    source_hash=source_hash,
                    locator_type="sheet_range",
                    sheet="Feedback_Explorer",
                    cell_range=f"A{row_number}:{get_column_letter(len(EXPECTED_HEADERS))}{row_number}",
                )
                annotations.append(
                    LegacyAnnotation(
                        annotation_id=stable_id(
                            "annotation", source_hash, source_record_id, str(row_number)
                        ),
                        feedback_record_id=record_id,
                        raw_feedback_type=_text(row[7]),
                        raw_level2_3=_text(row[9]),
                        raw_level2=_text(row[10]),
                        raw_level3=_text(row[11]),
                        row_locator=locator,
                    )
                )

        return FeedbackIngestionResult(
            source=source,
            input_annotation_rows=input_rows,
            input_unique_ids=len(grouped),
            records=records,
            annotations=annotations,
            transformations=transformations,
            quarantined=quarantined,
        )

"""Compute descriptive diagnostics only; labels are not model evaluation truth."""
import json
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
results = {}
for model in ('F70', 'G60', 'G70'):
    book = load_workbook(ROOT / f'data/raw/provided/{model}_feedback_hackathon.xlsx', read_only=True, data_only=True)
    sheet = book['Feedback_Explorer']
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows.pop(0)
    groups = defaultdict(list)
    for excel_row, row in enumerate(rows, 2):
        groups[row[0]].append((excel_row, row))
    first = [items[0][1] for items in groups.values()]
    text_groups = defaultdict(set)
    for row in first:
        text_groups[str(row[8]).strip().casefold()].add(row[0])
    results[model] = {
        'rows_excluding_header': len(rows), 'unique_ids': len(groups),
        'extra_annotation_rows': len(rows)-len(groups),
        'exact_duplicate_rows': len(rows)-len(set(rows)),
        'ids_with_conflicting_text': [key for key, items in groups.items() if len({r[8] for _, r in items}) > 1],
        'ids_with_conflicting_metadata': [key for key, items in groups.items() if len({r[1:7] for _, r in items}) > 1],
        'same_normalized_text_across_ids_groups': sum(len(ids)>1 for ids in text_groups.values()),
        'country_by_unique_id': dict(Counter(r[4] for r in first)),
        'source_by_unique_id': dict(Counter(r[1] for r in first)),
        'powertrain_by_unique_id': dict(Counter(r[6] for r in first)),
        'missing_feedback_type_rows': sum(r[7] is None for r in rows),
        'blank_text_rows': sum(not str(r[8] or '').strip() for r in rows),
        'unclassified_label_rows': sum(r[10] is None or r[10]=='no_class_found' for r in rows),
        'label_topics_unique_ids_top20': Counter(),
    }
    topics = defaultdict(set)
    for r in rows:
        if r[10] not in (None, 'no_class_found'):
            topics[r[10]].add(r[0])
    results[model]['label_topics_unique_ids_top20'] = sorted([(t,len(ids)) for t,ids in topics.items()],key=lambda x:-x[1])[:20]
    book.close()
(ROOT/'research/local/feedback_profile.json').write_text(json.dumps(results, indent=2),encoding='utf-8')
for model, data in results.items():
    print(model, json.dumps({k:v for k,v in data.items() if not k.endswith('top20')}))
    print('Existing-label topic counts:',data['label_topics_unique_ids_top20'][:8])
book=load_workbook(ROOT/'data/raw/provided/F70_G60_G68_G70_customer_studies.xlsx',data_only=True,read_only=True)
for sheet in book:
    print('STUDY:',sheet.title)
    for i,r in enumerate(sheet.iter_rows(values_only=True),1):
        if (sheet.title=='CN_EU_2025' and any(v is not None for v in r)) or (sheet.title=='US_2025' and (i<=25 or any(v is not None for v in r[4:]))):
            print(i, r)
book.close()

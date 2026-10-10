"""Read supplied BMW files locally; preserve bytes and produce a reproducible profile.

Usage: python scripts/inspect_supplied_data.py --zip PATH --briefing PATH
Requires openpyxl and pypdf. Does not run macros or upload any supplied data.
"""
import argparse
import hashlib
import json
import shutil
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--zip', required=True)
    parser.add_argument('--briefing', required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    raw = root / 'data/raw/provided'
    out = root / 'research/local'
    raw.mkdir(parents=True, exist_ok=True)
    out.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(args.zip) as archive:
        if sum(i.file_size for i in archive.infolist()) > 100_000_000:
            raise ValueError('Archive exceeds inspection size limit')
        for entry in archive.infolist():
            target = (raw / entry.filename).resolve()
            if not target.is_relative_to(raw.resolve()):
                raise ValueError('Archive path escapes raw data directory')
            if entry.is_dir():
                continue
            if target.suffix.lower() not in {'.xlsx', '.pdf'}:
                raise ValueError('Unexpected file type: ' + entry.filename)
            payload = archive.read(entry)
            if target.exists() and target.read_bytes() != payload:
                raise ValueError('Refusing to replace changed raw file: ' + str(target))
            target.parent.mkdir(parents=True, exist_ok=True)
            if not target.exists():
                target.write_bytes(payload)
    briefing = raw / Path(args.briefing).name
    if briefing.exists() and digest(briefing) != digest(Path(args.briefing)):
        raise ValueError('Refusing to overwrite changed briefing')
    if not briefing.exists():
        shutil.copyfile(args.briefing, briefing)
    profile = {'archive_sha256': digest(Path(args.zip)), 'files': []}
    for path in sorted(raw.iterdir()):
        item = {'file': path.name, 'bytes': path.stat().st_size, 'sha256': digest(path)}
        if path.suffix.lower() == '.pdf':
            pdf = PdfReader(path)
            item['pages'] = len(pdf.pages)
            text = '\n\n'.join(f'--- PAGE {n+1} ---\n{p.extract_text() or ""}' for n, p in enumerate(pdf.pages))
            (out / (path.stem + '.txt')).write_text(text, encoding='utf-8')
        elif path.suffix.lower() == '.xlsx':
            workbook = load_workbook(path, read_only=True, data_only=True)
            item['sheets'] = []
            for sheet in workbook:
                rows = list(sheet.iter_rows(values_only=True))
                nonempty = [(i + 1, r) for i, r in enumerate(rows) if any(v is not None for v in r)]
                cols = []
                for c in range(sheet.max_column):
                    values = [r[c] if c < len(r) else None for _, r in nonempty]
                    counts = Counter(str(v) for v in values if v is not None)
                    cols.append({'index_1based': c+1, 'nonempty': sum(v is not None for v in values),
                                 'distinct': len(counts), 'top_values': counts.most_common(12)})
                item['sheets'].append({'name': sheet.title, 'max_row': sheet.max_row,
                    'max_column': sheet.max_column, 'nonempty_rows': len(nonempty),
                    'first_rows': [{'excel_row': i, 'values': r} for i, r in nonempty[:12]],
                    'columns': cols})
            workbook.close()
        profile['files'].append(item)
    (out / 'profile.json').write_text(json.dumps(profile, indent=2, ensure_ascii=False, default=str), encoding='utf-8')
    print(json.dumps([{'file': f['file'], 'pages': f.get('pages'), 'sheets': [
        {k:s[k] for k in ('name','max_row','max_column','nonempty_rows')} for s in f.get('sheets',[])]} for f in profile['files']], indent=2))


if __name__ == '__main__':
    main()

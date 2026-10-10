"""Validate this preparation pack, not the unimplemented AI application.

Checks links, source hashes, observed data counts, study distributions and the
illustrative scoring arithmetic. Requires openpyxl. Run from any directory.
"""
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
checks = []
GENERATED_DIRS = {'.git', '.venv', '.uv-cache', '.worktrees', '.tmp', '.pytest_cache'}


def check(name, passed, detail):
    checks.append({'check': name, 'passed': bool(passed), 'detail': detail})


def main():
    markdown = sorted(
        path for path in ROOT.rglob('*.md')
        if not GENERATED_DIRS.intersection(path.relative_to(ROOT).parts)
    )
    broken = []
    for path in markdown:
        text = path.read_text(encoding='utf-8')
        for target in re.findall(r'\]\(([^)]+)\)', text):
            if target.startswith(('http://', 'https://', '#')):
                continue
            local = unquote(target.split('#', 1)[0]).strip('<>')
            if not (path.parent / local).exists():
                broken.append([str(path.relative_to(ROOT)), local])
    check('markdown_local_links', not broken, broken)

    profile = json.loads((ROOT/'research/local/profile.json').read_text(encoding='utf-8'))
    mismatches = []
    for item in profile['files']:
        raw = ROOT/'data/raw/provided'/item['file']
        if hashlib.sha256(raw.read_bytes()).hexdigest() != item['sha256']:
            mismatches.append(item['file'])
    check('raw_source_hashes', not mismatches, {'files': len(profile['files']), 'mismatches': mismatches})

    counts = {}
    for model, expected in {'F70':(4076,2128),'G60':(5005,4000),'G70':(4690,4000)}.items():
        book = load_workbook(ROOT/f'data/raw/provided/{model}_feedback_hackathon.xlsx', read_only=True, data_only=True)
        records = list(book['Feedback_Explorer'].iter_rows(values_only=True))[1:]
        observed = (len(records), len({r[0] for r in records}))
        counts[model] = {'rows': observed[0], 'unique_ids': observed[1]}
        check('feedback_counts_'+model, observed == expected, counts[model])
        if model == 'G60':
            expected_ids = {24:'G60-0021',59:'G60-0051',95:'G60-0078',103:'G60-0083',114:'G60-0092',287:'G60-0236',344:'G60-0271'}
            bad = {row:records[row-2][0] for row,record_id in expected_ids.items() if records[row-2][0] != record_id}
            check('demo_source_row_ids', not bad, {'verified':len(expected_ids),'mismatches':bad})
        book.close()

    book = load_workbook(ROOT/'data/raw/provided/sales_volumes.xlsx',read_only=True,data_only=True)
    sales = {}
    for sheet in book:
        rows = list(sheet.iter_rows(values_only=True))[1:]
        sales[sheet.title] = [sum(r[col] for r in rows) for col in (2,3,4)]
    check('sales_totals', sales == {'F70':[105000,113000,118000],'G60_G68':[278600,296000,316000],'G70':[34000,42000,52000]},sales)
    book.close()

    book = load_workbook(ROOT/'data/raw/provided/F70_G60_G68_G70_customer_studies.xlsx',read_only=True,data_only=True)
    cn_eu = list(book['CN_EU_2025'].iter_rows(values_only=True))
    mean_values = [value for row in cn_eu if row[0]=='Mean' for value in row[1:6] if isinstance(value,(int,float))]
    check('study_cn_eu_means',len(mean_values)==185,{'mean_cells':len(mean_values)})
    labels = {'I Hate It','A Failure','Unsatisfactory','Satisfactory','Excellent','Delightful','I Love It'}
    blocks = []
    current = None
    for row in book['US_2025'].iter_rows(values_only=True):
        if row[0] is not None and row[0] not in labels and row[0]!='Sample total':
            current = {'category': row[0], 'rows': []}
            blocks.append(current)
        elif row[0] in labels:
            if current is None:
                raise ValueError('Response precedes study category')
            current['rows'].append(row)
    errors = []
    for block in blocks:
        actual_labels = Counter(row[0] for row in block['rows'])
        if actual_labels != Counter({label:1 for label in labels}):
            errors.append([block['category'],'response_labels'])
        for col in (2,3):
            values = [row[col] for row in block['rows']]
            if not all(isinstance(v,(int,float)) and 0<=v<=1 for v in values) or abs(sum(values)-1)>0.005:
                errors.append([block['category'],col+1,'invalid_distribution'])
    check('study_us_distributions',len(blocks)==103 and not errors,{'categories':len(blocks),'model_distributions':len(blocks)*2,'errors':errors})
    book.close()

    def value(xs, weights):
        return 100*sum(x*w/4 for x,w in zip(xs,weights))
    values = [value((3,2,3,1),(.25,)*4), value((2,3,2,3),(.25,)*4),value((3,2,3,1),(.4,.1,.4,.1)),value((2,3,2,3),(.4,.1,.4,.1))]
    check('illustrative_score_arithmetic',all(abs(a-b)<1e-9 for a,b in zip(values,[56.25,62.5,67.5,55])),values)

    seeds = [json.loads(line) for line in (ROOT/'research/web_evidence_seed.jsonl').read_text(encoding='utf-8').splitlines() if line.strip()]
    check('web_seeds_are_labelled_curated',len(seeds)==5 and all(s['needs_refetch_and_locator_validation'] and s['source_snapshot_hash'] is None for s in seeds),{'count':len(seeds)})
    agent_file = ROOT/'internal/AGENTS.md'
    if agent_file.exists():
        check('internal_agent_file_small', agent_file.stat().st_size<10000, {'bytes':agent_file.stat().st_size})

    result = {'scope':'Preparation artifacts only; no application/model quality tested',
              'markdown_files':len(markdown), 'markdown_words':sum(len(p.read_text(encoding='utf-8').split()) for p in markdown),
              'passed':all(c['passed'] for c in checks), 'checks':checks}
    (ROOT/'research/local/preparation_validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))
    if not result['passed']:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
